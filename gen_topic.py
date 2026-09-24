from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path
from dotenv import load_dotenv
from datetime import datetime, timezone
import json
import os
import time
import random

load_dotenv()

try:
    import db
    MONGO_AVAILABLE = True
except Exception as e:
    print(f"[db] MongoDB not available: {e}")
    MONGO_AVAILABLE = False

GROQ_API_KEY = os.getenv("GROQ_API_KEY")
GROQ_API_KEY_2 = os.getenv("GROQ_API_KEY_2")
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")
MISTRAL_API_KEY = os.getenv("MISTRAL_API_KEY")

USAGE_FILE = Path("provider_usage.json")

LIMITS = {
    "groq_1": {"tokens_per_day": 200_000},
    "groq_2": {"tokens_per_day": 200_000},
    "gemini": {"requests_per_day": 1500, "requests_per_minute": 15},
    "mistral": {"tokens_per_day": 5_000_000},
}

DEFAULT_USAGE = {
    "date_utc": "",
    "groq_1": {"tokens_today": 0},
    "groq_2": {"tokens_today": 0},
    "gemini": {"requests_today": 0, "last_minute_requests": []},
    "mistral": {"tokens_today": 0},
}


class RateLimitReached(Exception):
    pass


def load_usage():
    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    data = dict(DEFAULT_USAGE)
    data["date_utc"] = today
    return data


def save_usage(data):
    with open(USAGE_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2)


def estimate_tokens(text):
    return max(1, len(text) // 4)


groq_clients = []
for k in [GROQ_API_KEY, GROQ_API_KEY_2]:
    if k:
        try:
            from groq import Groq
            groq_clients.append(Groq(api_key=k))
        except Exception as e:
            print(f"[groq init failed] {e}")

# Gemini disabled - package removed to fix Render build
gemini_model = None
print("[gemini] Skipped (not installed)")

mistral_client = None
if MISTRAL_API_KEY:
    try:
        from mistralai import Mistral
        mistral_client = Mistral(api_key=MISTRAL_API_KEY)
    except Exception as e:
        print(f"[mistral init failed] {e}")


UNIVERSITY_NAME = "Metropolitan International University"
FOOTER_TEXT = "www.miu.ac.ug  |  info@miu.ac.ug  |  +256 772 561 957  |  Kampala - Mbarara - Kisoro Campuses"

PRIMARY_COLOR = RGBColor(0, 130, 60)
ACCENT_COLOR = RGBColor(200, 30, 40)
WHITE = RGBColor(255, 255, 255)
DARK_TEXT = RGBColor(30, 30, 30)

LOGO_PATH = "miu_logo.png"

BASE_DIR = Path(__file__).parent
CACHE_ROOT = BASE_DIR / "generated_content"
SLIDES_ROOT = BASE_DIR / "output_slides"
LOGS_DIR = BASE_DIR / "logs"

for d in (CACHE_ROOT, SLIDES_ROOT, LOGS_DIR):
    d.mkdir(exist_ok=True)


# ============================================================
# PATH BUILDERS
# ============================================================

def clean_name(text):
    return "".join(c if c.isalnum() else "_" for c in text).strip("_")


def programme_dir(program_code):
    return clean_name(program_code)


def year_dir(year):
    return f"Year_{year}"


def semester_dir(semester):
    return f"Semester_{semester}"


def course_dir(course_code, course_name):
    return f"{clean_name(course_code)}_{clean_name(course_name)}"


def topic_filename(topic_number, topic_name):
    return f"{topic_number:02d}_{clean_name(topic_name)}"


def build_paths(program_code, year, semester, course_code, course_name,
                topic_number, topic_name):
    course_folder = course_dir(course_code, course_name)
    stem = topic_filename(topic_number, topic_name)

    cache_dir = (CACHE_ROOT / programme_dir(program_code)
                 / year_dir(year) / semester_dir(semester) / course_folder)
    slides_dir = (SLIDES_ROOT / programme_dir(program_code)
                  / year_dir(year) / semester_dir(semester) / course_folder)

    cache_dir.mkdir(parents=True, exist_ok=True)
    slides_dir.mkdir(parents=True, exist_ok=True)

    return (cache_dir, cache_dir / f"{stem}.json",
            slides_dir, slides_dir / f"{stem}.pptx")


def log_progress(message):
    LOGS_DIR.mkdir(exist_ok=True)
    log_file = LOGS_DIR / "progress.log"
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with open(log_file, "a", encoding="utf-8") as f:
        f.write(f"[{timestamp}] {message}\n")


# ============================================================
# PROVIDER LOGIC
# ============================================================

def pick_provider(usage, estimated_tokens):
    if mistral_client:
        m = usage["mistral"]
        if m["tokens_today"] < LIMITS["mistral"]["tokens_per_day"]:
            return "mistral", None

    groq_choices = []
    for i, key_name in enumerate(["groq_1", "groq_2"], 1):
        if i > len(groq_clients):
            continue
        used = usage[key_name]["tokens_today"]
        remaining = LIMITS[key_name]["tokens_per_day"] - used
        if remaining >= estimated_tokens:
            groq_choices.append((remaining, key_name, i))

    if groq_choices:
        groq_choices.sort(reverse=True)
        return groq_choices[0][1], groq_choices[0][2]

    if gemini_model:
        g = usage["gemini"]
        if g["requests_today"] < LIMITS["gemini"]["requests_per_day"]:
            now = time.time()
            recent = [t for t in g["last_minute_requests"] if now - t < 60]
            if len(recent) < LIMITS["gemini"]["requests_per_minute"]:
                return "gemini", None

    return None, None


def call_groq(client_index, system_prompt, user_prompt):
    client = groq_clients[client_index - 1]
    print(f"    [content] Calling Groq key #{client_index}...")
    response = client.chat.completions.create(
        model="openai/gpt-oss-120b",
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt}
        ],
        temperature=0.7
    )
    return response.choices[0].message.content


def call_gemini(system_prompt, user_prompt):
    # Gemini is disabled - package removed
    raise Exception("Gemini disabled")


def call_mistral(system_prompt, user_prompt, max_retries=5):
    for attempt in range(max_retries):
        try:
            print(f"    [content] Calling Mistral... (attempt {attempt + 1})")
            response = mistral_client.chat.complete(
                model="mistral-small-latest",
                messages=[
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": user_prompt}
                ]
            )
            return response.choices[0].message.content
        except Exception as e:
            err = str(e)
            is_rate_limit = ("429" in err or "rate" in err.lower()
                             or "1300" in err or "rate_limited" in err.lower())
            if is_rate_limit and attempt < max_retries - 1:
                wait = (2 ** attempt) + random.uniform(0, 1)
                print(f"    [mistral] Rate limit hit. Waiting {wait:.1f}s...")
                time.sleep(wait)
            else:
                raise


def _try_other_groq(current_provider, usage, system_prompt, user_prompt,
                    estimated_tokens):
    for i, key_name in enumerate(["groq_1", "groq_2"], 1):
        if i > len(groq_clients) or key_name == current_provider:
            continue
        if usage[key_name]["tokens_today"] >= LIMITS[key_name]["tokens_per_day"]:
            continue
        try:
            print(f"    [fallback] Trying Groq key #{i}...")
            raw = call_groq(i, system_prompt, user_prompt)
            usage[key_name]["tokens_today"] += estimated_tokens
            save_usage(usage)
            return raw
        except Exception as e:
            err = str(e)
            print(f"    [{key_name} FAILED] {err[:200]}")
            if "429" in err or "rate_limit" in err.lower():
                usage[key_name]["tokens_today"] += 200_000
                save_usage(usage)
    return None


# ============================================================
# PROMPTS
# ============================================================

def build_prompts(topic, course_name, course_code, topic_number, total_topics):
    system_prompt = (
        "You are a friendly university lecturer writing teaching slides "
        "for undergraduate students learning this topic for the first time. "
        "Make the content CLEAR, SIMPLE and EASY TO UNDERSTAND. "
        "You MUST return ONLY valid JSON - no markdown, no explanation. "
        'The JSON must be exactly this shape: '
        '{"slides": [{"title": "...", "bullets": ["...", "..."]}]} '
        "Each slide object must have: "
        '"title" (short slide heading) and '
        '"bullets" (list of 8 to 10 full explanatory sentences).'
    )
    user_prompt = (
        f"Course: {course_name} ({course_code})\n"
        f"Topic {topic_number} of {total_topics}: {topic}\n\n"
        f"Write 15 teaching slides using this structure:\n"
        f"1. Learning Objectives\n"
        f"2. Introduction - What is {topic}?\n"
        f"3. Why {topic} Matters\n"
        f"4. Key Ideas - Part 1\n"
        f"5. Key Ideas - Part 2\n"
        f"6. How It Works - Step by Step\n"
        f"7. Important Terms Explained\n"
        f"8. Real-Life Examples\n"
        f"9. Worked Example\n"
        f"10. Common Mistakes and How to Avoid Them\n"
        f"11. Advantages and Limitations\n"
        f"12. Best Practices\n"
        f"13. Modern Developments\n"
        f"14. Summary - What You Should Remember\n"
        f"15. Practice Questions and Further Reading\n\n"
        f"Every bullet must be a full sentence or two. Write in simple English. "
        f"Define every technical term the first time it appears.\n\n"
        f"Return only the JSON object."
    )
    return system_prompt, user_prompt


def parse_json_robust(raw):
    if not raw:
        return None
    raw = raw.strip()
    if raw.startswith("```"):
        parts = raw.split("```")
        if len(parts) >= 2:
            raw = parts[1]
            if raw.startswith("json"):
                raw = raw[4:]
    raw = raw.strip()
    start = raw.find("{")
    end = raw.rfind("}")
    if start != -1 and end != -1:
        raw = raw[start:end + 1]
    try:
        return json.loads(raw)
    except Exception as e:
        print(f"    [parse attempt 1 failed] {str(e)[:100]}")
    try:
        fixed = raw.replace(",}", "}").replace(",]", "]")
        fixed = "".join(ch for ch in fixed if ch >= " " or ch in "\n\r\t")
        return json.loads(fixed)
    except Exception as e:
        print(f"    [parse attempt 2 failed] {str(e)[:100]}")
    try:
        from json_repair import repair_json
        fixed = repair_json(raw)
        result = json.loads(fixed)
        print("    [parse attempt 3 succeeded]")
        return result
    except Exception as e:
        print(f"    [parse attempt 3 failed] {str(e)[:100]}")
    return None


# ============================================================
# CONTENT GENERATION
# ============================================================

def generate_slide_content(topic, course_name, course_code, topic_number,
                           total_topics, cache_file):
    if cache_file.exists():
        with open(cache_file, "r", encoding="utf-8") as f:
            return json.load(f)

    print(f"    [content] Generating: '{topic}'")
    system_prompt, user_prompt = build_prompts(
        topic, course_name, course_code, topic_number, total_topics
    )

    total_prompt = system_prompt + user_prompt
    estimated_tokens = estimate_tokens(total_prompt) + 3500

    usage = load_usage()
    provider, groq_index = pick_provider(usage, estimated_tokens)
    if provider is None:
        print("    [RATE LIMIT] All providers exhausted.")
        raise RateLimitReached("All providers out of quota")

    raw = None

    if provider == "mistral":
        try:
            raw = call_mistral(system_prompt, user_prompt)
            usage["mistral"]["tokens_today"] += estimated_tokens
            save_usage(usage)
        except Exception as e:
            err = str(e)
            print(f"    [mistral FAILED] {err[:200]}")
            if "429" in err or "rate" in err.lower() or "quota" in err.lower():
                usage["mistral"]["tokens_today"] += 5_000_000
                save_usage(usage)

    elif provider in ("groq_1", "groq_2"):
        try:
            raw = call_groq(groq_index, system_prompt, user_prompt)
            usage[provider]["tokens_today"] += estimated_tokens
            save_usage(usage)
        except Exception as e:
            err = str(e)
            print(f"    [{provider} FAILED] {err[:200]}")
            if "429" in err or "rate_limit" in err.lower() or "quota" in err.lower():
                usage[provider]["tokens_today"] += 200_000
                save_usage(usage)

    elif provider == "gemini":
        try:
            raw = call_gemini(system_prompt, user_prompt)
            usage["gemini"]["requests_today"] += 1
            usage["gemini"]["last_minute_requests"].append(time.time())
            now = time.time()
            usage["gemini"]["last_minute_requests"] = [
                t for t in usage["gemini"]["last_minute_requests"] if now - t < 60
            ]
            save_usage(usage)
        except Exception as e:
            err = str(e)
            print(f"    [gemini FAILED] {err[:200]}")
            if "429" in err or "quota" in err.lower():
                usage["gemini"]["requests_today"] += 1500
                save_usage(usage)

    if not raw and provider != "mistral" and mistral_client:
        try:
            print("    [fallback] Trying Mistral...")
            raw = call_mistral(system_prompt, user_prompt)
            usage["mistral"]["tokens_today"] += estimated_tokens
            save_usage(usage)
        except Exception as e:
            print(f"    [mistral FAILED] {str(e)[:200]}")

    if not raw:
        raw = _try_other_groq(provider, usage, system_prompt,
                              user_prompt, estimated_tokens)

    if not raw and provider != "gemini" and gemini_model:
        try:
            print("    [fallback] Trying Gemini...")
            raw = call_gemini(system_prompt, user_prompt)
            usage["gemini"]["requests_today"] += 1
            usage["gemini"]["last_minute_requests"].append(time.time())
            save_usage(usage)
        except Exception as e:
            print(f"    [gemini FAILED] {str(e)[:200]}")

    if not raw:
        print("    [RATE LIMIT] No provider produced content.")
        raise RateLimitReached("All providers exhausted")

    result = parse_json_robust(raw)
    if result is None:
        print("    [content FAILED] Could not parse JSON")
        raise RateLimitReached("JSON unparseable")

    cache_file.parent.mkdir(parents=True, exist_ok=True)
    with open(cache_file, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    return result


# ============================================================
# SLIDE BUILDING
# ============================================================

def add_logo_badge(slide, prs, center_x, top, size):
    if not os.path.exists(LOGO_PATH):
        return
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL,
                                    center_x - size / 2, top, size, size)
    circle.fill.solid()
    circle.fill.fore_color.rgb = WHITE
    circle.line.color.rgb = PRIMARY_COLOR
    circle.line.width = Pt(2.5)
    padding = Emu(int(size * 0.08))
    try:
        slide.shapes.add_picture(
            LOGO_PATH,
            center_x - size / 2 + padding,
            top + padding,
            width=size - padding * 2,
            height=size - padding * 2
        )
    except Exception as e:
        print(f"    [logo failed] {e}")


def add_footer(slide, prs):
    if os.path.exists(LOGO_PATH):
        try:
            slide.shapes.add_picture(
                LOGO_PATH, Inches(0.35),
                prs.slide_height - Inches(0.55), height=Inches(0.42)
            )
        except Exception:
            pass
    footer = slide.shapes.add_textbox(
        Inches(1.2), prs.slide_height - Inches(0.45),
        prs.slide_width - Inches(1.5), Inches(0.35)
    )
    p = footer.text_frame.paragraphs[0]
    p.text = FOOTER_TEXT
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(80, 80, 80)
    p.alignment = PP_ALIGN.RIGHT


def add_header_bar(slide, prs, title_text, topic_label):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
        prs.slide_width, Inches(1.05)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = PRIMARY_COLOR
    shape.line.fill.background()

    box = slide.shapes.add_textbox(Inches(0.4), Inches(0.22),
                                   Inches(8.5), Inches(0.65))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = WHITE

    tag = slide.shapes.add_textbox(Inches(9.0), Inches(0.30),
                                   Inches(4.0), Inches(0.5))
    tf2 = tag.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = topic_label
    p2.font.size = Pt(11)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.alignment = PP_ALIGN.RIGHT

    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.05),
        prs.slide_width, Inches(0.06)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = ACCENT_COLOR
    accent.line.fill.background()


def build_title_slide(prs, course_name, course_code, topic,
                      topic_number, total_topics, year, semester):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
                                prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = PRIMARY_COLOR
    bg.line.fill.background()

    add_logo_badge(slide, prs, prs.slide_width / 2, Inches(0.7), Inches(1.6))

    box = slide.shapes.add_textbox(Inches(0.8), Inches(2.5),
                                   prs.slide_width - Inches(1.6), Inches(2.4))
    tf = box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = UNIVERSITY_NAME
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER

    p2 = tf.add_paragraph()
    p2.text = f"{course_name} ({course_code})"
    p2.font.size = Pt(16)
    p2.font.color.rgb = WHITE
    p2.alignment = PP_ALIGN.CENTER

    p3 = tf.add_paragraph()
    p3.text = f"Year {year} - Semester {semester}"
    p3.font.size = Pt(13)
    p3.font.color.rgb = WHITE
    p3.alignment = PP_ALIGN.CENTER

    p4 = tf.add_paragraph()
    p4.text = topic
    p4.font.size = Pt(30)
    p4.font.bold = True
    p4.font.color.rgb = WHITE
    p4.alignment = PP_ALIGN.CENTER

    p5 = tf.add_paragraph()
    p5.text = f"Topic {topic_number} of {total_topics}"
    p5.font.size = Pt(14)
    p5.font.color.rgb = WHITE
    p5.alignment = PP_ALIGN.CENTER

    add_footer(slide, prs)
    return slide


def build_content_slide(prs, slide_data, course_code, topic, slide_no, total_slides):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_bar(
        slide, prs,
        title_text=slide_data.get("title", f"Slide {slide_no}"),
        topic_label=f"{course_code} - {topic}"
    )
    body = slide.shapes.add_textbox(
        Inches(0.6), Inches(1.35),
        prs.slide_width - Inches(1.2),
        prs.slide_height - Inches(2.0)
    )
    tf = body.text_frame
    tf.word_wrap = True

    bullets = slide_data.get("bullets", []) or ["No content available."]
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"- {bullet}"
        p.font.size = Pt(13)
        p.font.color.rgb = DARK_TEXT
        p.space_after = Pt(6)

    page_no = slide.shapes.add_textbox(
        prs.slide_width - Inches(1.3),
        prs.slide_height - Inches(0.85),
        Inches(1.0), Inches(0.3)
    )
    pp = page_no.text_frame.paragraphs[0]
    pp.text = f"{slide_no} / {total_slides}"
    pp.font.size = Pt(9)
    pp.font.color.rgb = RGBColor(120, 120, 120)
    pp.alignment = PP_ALIGN.RIGHT

    add_footer(slide, prs)
    return slide


# ============================================================
# DECK BUILDER
# ============================================================

def build_topic_deck(program_code, year, semester,
                     course_name, course_code,
                     topic, topic_number, total_topics):
    cache_dir, cache_file, slides_dir, slides_file = build_paths(
        program_code, year, semester,
        course_code, course_name,
        topic_number, topic
    )

    # If pptx exists on disk AND in MongoDB, skip
    if slides_file.exists() and MONGO_AVAILABLE and db.pptx_exists_in_mongo(course_code, topic_number):
        print(f"    [skip] Already exists (disk + MongoDB)")
        return slides_file

    # If pptx exists on disk but not MongoDB, upload it
    if slides_file.exists() and MONGO_AVAILABLE and not db.pptx_exists_in_mongo(course_code, topic_number):
        try:
            file_id = db.save_pptx_to_mongo(course_code, topic_number, slides_file.name, slides_file)
            if file_id:
                print(f"    [gridfs] Backed up existing pptx to MongoDB")
        except Exception as e:
            print(f"    [gridfs] backup failed: {e}")
        return slides_file

    content = generate_slide_content(
        topic, course_name, course_code,
        topic_number, total_topics, cache_file
    )
    slides_data = content.get("slides", [])
    if not slides_data:
        raise RateLimitReached("No slides returned")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    build_title_slide(prs, course_name, course_code,
                      topic, topic_number, total_topics, year, semester)

    total = len(slides_data)
    for i, slide_data in enumerate(slides_data, 1):
        build_content_slide(prs, slide_data, course_code, topic, i, total)

    prs.save(str(slides_file))
    rel = slides_file.relative_to(BASE_DIR)
    print(f"    [saved] {rel}")
    log_progress(f"[Y{year}S{semester}] [{course_code}] {topic} - DONE ({rel})")

    # Save to MongoDB
    if MONGO_AVAILABLE:
        try:
            # 1. Save pptx BYTES to GridFS
            file_id = db.save_pptx_to_mongo(
                course_code, topic_number, slides_file.name, slides_file
            )
            if file_id:
                print(f"    [gridfs] Saved to MongoDB (id={file_id[:8]}...)")

            # 2. Archive previous content version
            existing = db.col_content().find_one(
                {"course_code": course_code, "topic_number": topic_number}
            )
            if existing and existing.get("content"):
                try:
                    history_col = db.get_db()["generated_content_history"]
                    history_col.insert_one({
                        "course_code": course_code,
                        "topic_number": topic_number,
                        "topic": topic,
                        "course_name": course_name,
                        "program": program_code,
                        "year": year,
                        "semester": semester,
                        "content": existing.get("content"),
                        "archived_at": datetime.now(timezone.utc),
                    })
                    print(f"    [history] Archived previous version")
                except Exception as he:
                    print(f"    [history] archive failed: {he}")

            # 3. Slide metadata
            db.col_slides().update_one(
                {"course_code": course_code, "topic_number": topic_number},
                {"$set": {
                    "program": program_code,
                    "year": year,
                    "semester": semester,
                    "course_name": course_name,
                    "topic": topic,
                    "filename": slides_file.name,
                    "rel_path": str(rel).replace("\\", "/"),
                    "size_bytes": slides_file.stat().st_size,
                    "gridfs_id": file_id,
                    "updated_at": datetime.now(timezone.utc),
                }},
                upsert=True,
            )

            # 4. Current content
            db.col_content().update_one(
                {"course_code": course_code, "topic_number": topic_number},
                {"$set": {
                    "program": program_code,
                    "year": year,
                    "semester": semester,
                    "course_code": course_code,
                    "course_name": course_name,
                    "topic_number": topic_number,
                    "topic": topic,
                    "content": content,
                    "updated_at": datetime.now(timezone.utc),
                }},
                upsert=True,
            )

            # 5. Activity log
            db.log_activity(
                "generate",
                f"{course_code} - {topic}",
                user="system"
            )
        except Exception as e:
            print(f"    [db] MongoDB save failed: {e}")

    return slides_file


# ============================================================
# MAIN
# ============================================================

def main():
    curriculum_file = BASE_DIR / "curriculum.json"

    with open(curriculum_file, "r", encoding="utf-8") as f:
        curriculum = json.load(f)

    for program in curriculum["programs"]:
        program_code = program["code"]

        for year_block in program["years"]:
            year = year_block["year"]

            for semester_block in year_block["semesters"]:
                semester = semester_block["semester"]

                for course in semester_block["courses"]:
                    course_name = course["name"]
                    course_code = course["code"]
                    topics = course.get("topics", [])
                    total_topics = len(topics)

                    print(f"\n=== {course_code} - {course_name} "
                          f"(Year {year}, Sem {semester}) ===")

                    for idx, topic in enumerate(topics, 1):
                        try:
                            build_topic_deck(
                                program_code, year, semester,
                                course_name, course_code,
                                topic, idx, total_topics
                            )
                        except RateLimitReached as e:
                            print(f"\n[STOP] {e}")
                            log_progress(f"[STOP] Rate limit - {course_code}/{topic}")
                            if MONGO_AVAILABLE:
                                try:
                                    db.log_activity("rate_limit", str(e), user="system")
                                except Exception:
                                    pass
                            print("Re-run later - cached topics will be skipped.")
                            return
                        except Exception as e:
                            print(f"[ERROR] {course_code} / {topic}: {e}")
                            log_progress(f"[ERROR] {course_code}/{topic}: {e}")
                            continue


if __name__ == "__main__":
    main()